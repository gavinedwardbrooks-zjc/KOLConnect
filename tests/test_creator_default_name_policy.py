from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
APP = ROOT / "app"
if str(APP) not in sys.path:
    sys.path.insert(0, str(APP))

from campaign_creator_repository import CampaignCreatorRepository
from campaign_repository import CampaignRepository
from creator_repository import CreatorRepository
from domain.creator_naming import default_creator_name
from product_repository import ProductRepository
from storage.sqlite_creator_repository import SQLiteCreatorRepository
from storage.sqlite_workbook_store import SQLiteWorkbookStore


def analysis(
    suffix: str,
    *,
    username: str = "",
    creator_name: str = "",
    profile_title: str = "",
    profile_url: str = "https://www.tiktok.com/@kaho.brazuka",
    creator_id: str = "",
) -> dict:
    task_id = f"task_20261022T120000Z_{suffix}"
    payload = {
        "schema_version": "1.0",
        "analysis_id": f"analysis_{task_id}",
        "task_id": task_id,
        "account_uid": f"tiktok|{profile_url}",
        "imported_at": "2026-10-22T12:00:00Z",
        "source": "chrome_extension",
        "creator": {
            "creator_name": creator_name,
            "username": username,
            "profile_title": profile_title,
            "platform": "TikTok",
            "profile_url": profile_url,
        },
        "field_observations": {
            "creator_name": {"observed": bool(creator_name)},
            "username": {"observed": bool(username)},
            "profile_url": {"observed": True},
            "platform": {"observed": True},
        },
        "video_analysis": {},
        "videos": [],
        "creator_insight": {},
    }
    if creator_id:
        payload["creator_id"] = creator_id
    return payload


class CreatorDefaultNamePolicyTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.workbook_path = self.root / "Creator_Library.xlsx"
        self.log_patch = mock.patch("creator_repository.log_event")
        self.log_patch.start()
        self.repository = CreatorRepository(self.workbook_path)

    def tearDown(self) -> None:
        self.log_patch.stop()
        self.temp_dir.cleanup()

    def test_default_rule_prefers_display_then_handle_then_profile_title(self) -> None:
        self.assertEqual(
            "显示名称",
            default_creator_name(
                display_name=" 显示名称 ", username="handle", profile_title="Profile title"
            ),
        )
        self.assertEqual(
            "@kaho.brazuka",
            default_creator_name(display_name="N/A", username="@kaho.brazuka", profile_title="Title"),
        )
        self.assertEqual(
            "Profile title",
            default_creator_name(display_name="-", username="undefined", profile_title="Profile title"),
        )
        self.assertEqual("", default_creator_name(display_name="null", username="@N/A"))

    def test_extension_add_seeds_name_from_display_or_handle(self) -> None:
        display = self.repository.saveCreator(
            analysis("aaaaaaaa", username="@display.handle", creator_name="Creator 名称"),
            extension_action="ADD",
        )
        handle = self.repository.saveCreator(
            analysis(
                "bbbbbbbb",
                username="@kaho.brazuka",
                profile_url="https://www.tiktok.com/@kaho.brazuka.second",
            ),
            extension_action="ADD",
        )
        missing = self.repository.saveCreator(
            analysis(
                "cccccccc",
                profile_url="https://www.tiktok.com/",
            ),
            extension_action="ADD",
        )

        self.assertEqual("Creator 名称", self.repository.getCreatorDetail(display["creator_id"])["record"]["creator_name"])
        self.assertEqual("@kaho.brazuka", self.repository.getCreatorDetail(handle["creator_id"])["record"]["creator_name"])
        self.assertEqual("", self.repository.getCreatorDetail(missing["creator_id"])["record"]["creator_name"])

    def test_existing_manual_name_wins_across_refresh_and_second_account(self) -> None:
        first = self.repository.saveCreator(
            analysis("dddddddd", username="first", creator_name="Initial"),
            extension_action="ADD",
        )
        creator_id = first["creator_id"]
        self.repository.updateCreator(creator_id, {"creator_name": "人工命名"})

        self.repository.saveCreator(
            analysis("eeeeeeee", username="refreshed", creator_name="Observed", creator_id=creator_id),
            extension_action="UPDATE",
        )
        self.repository.saveCreator(
            analysis(
                "ffffffff",
                username="second.account",
                creator_name="Second Platform Name",
                profile_url="https://www.tiktok.com/@second.account",
                creator_id=creator_id,
            )
        )

        detail = self.repository.getCreatorDetail(creator_id)
        self.assertEqual("人工命名", detail["record"]["creator_name"])
        self.assertEqual(2, len(detail["accounts"]))

    def test_task_import_seeds_blank_name_but_never_overwrites_existing_name(self) -> None:
        record = {
            "account_uid": "tiktok|https://www.tiktok.com/@task.handle",
            "platform": "TikTok",
            "profile_url": "https://www.tiktok.com/@task.handle",
            "creator_name": "",
            "scrape_status": "success",
        }
        first = self.repository.importTaskResults(
            "task_20261022T130000Z_aaaaaaaa",
            [record],
            source="系统抓取",
            imported_at="2026-10-22T13:00:00Z",
        )
        creator_id = first["creator_ids"][0]
        self.assertEqual("@task.handle", self.repository.getCreatorDetail(creator_id)["record"]["creator_name"])

        self.repository.updateCreator(creator_id, {"creator_name": "人工保留"})
        self.repository.importTaskResults(
            "task_20261022T140000Z_bbbbbbbb",
            [{**record, "creator_name": "Later Scraped Name"}],
            source="系统抓取",
            imported_at="2026-10-22T14:00:00Z",
        )
        self.assertEqual("人工保留", self.repository.getCreatorDetail(creator_id)["record"]["creator_name"])

    def test_account_link_repairs_only_a_truly_blank_existing_creator(self) -> None:
        database = self.root / "kolconnect.db"
        store = SQLiteWorkbookStore.initialize_empty(database, reference="creator-name-test")
        repository = SQLiteCreatorRepository(store)
        now = "2026-10-22T15:00:00Z"
        with store.factory.write_transaction() as connection:
            connection.execute(
                "INSERT INTO creators(creator_id,name,status,created_at,updated_at) VALUES (?,?,?,?,?)",
                ("blank", "", "discovered", now, now),
            )
            connection.execute(
                "INSERT INTO creators(creator_id,name,status,created_at,updated_at) VALUES (?,?,?,?,?)",
                ("literal", "未命名达人", "discovered", now, now),
            )

        blank_url = "https://www.tiktok.com/@linked.handle"
        literal_url = "https://www.tiktok.com/@literal.handle"
        repository.linkNewAccountToCreator(
            "blank", "TikTok", blank_url, repository._build_account_uid("TikTok", blank_url)
        )
        repository.linkNewAccountToCreator(
            "literal", "TikTok", literal_url, repository._build_account_uid("TikTok", literal_url)
        )

        self.assertEqual("@linked.handle", repository.getCreatorDetail("blank")["record"]["creator_name"])
        self.assertEqual("未命名达人", repository.getCreatorDetail("literal")["record"]["creator_name"])

    def test_campaign_assignment_repairs_and_displays_handle(self) -> None:
        saved = self.repository.saveCreator(
            analysis("11111111", username="campaign.handle", creator_name="Seed"),
            extension_action="ADD",
        )
        creator_id = saved["creator_id"]
        workbook = load_workbook(self.workbook_path)
        headers = [str(cell.value or "") for cell in workbook["Creators"][1]]
        name_column = headers.index("name") + 1
        for row in range(2, workbook["Creators"].max_row + 1):
            if workbook["Creators"].cell(row, 1).value == creator_id:
                workbook["Creators"].cell(row, name_column).value = ""
        workbook.save(self.workbook_path)
        workbook.close()

        product = ProductRepository(self.workbook_path).createProduct({"name": "Product"})
        campaign = CampaignRepository(self.workbook_path).createCampaign(
            {"name": "Campaign", "product_id": product["product_id"], "platform": "TikTok"}
        )
        account = self.repository.getCreatorAccounts(creator_id)[0]
        relation = CampaignCreatorRepository(self.workbook_path).createCampaignCreator(
            {
                "campaign_id": campaign["campaign_id"],
                "creator_id": creator_id,
                "account_id": account["account_id"],
            }
        )

        self.assertEqual("@campaign.handle", relation["creator_name"])
        self.assertEqual("@campaign.handle", self.repository.getCreatorDetail(creator_id)["record"]["creator_name"])


if __name__ == "__main__":
    unittest.main()
