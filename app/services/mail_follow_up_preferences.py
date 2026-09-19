"""Persistent, group-local user decisions for the derived Mail follow-up queue."""

from __future__ import annotations

import csv
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook

from services.mail_follow_up import RELIABLE_MATCH_STATUSES, _parse_utc, derive_follow_up_queue
from services.mail_inbox_facts import normalize_address


EXPORT_HEADERS = (
    "creator_id", "creator_name", "correspondent_email", "waiting_for", "days_waiting",
    "last_inbound_at", "last_outbound_at", "last_mail_at", "synced_inbound_count",
    "synced_outbound_count", "partial_history",
)


def _utc_now(now: datetime | None) -> datetime:
    return (now or datetime.now(timezone.utc)).astimezone(timezone.utc).replace(microsecond=0)


def _utc_text(value: datetime) -> str:
    return value.isoformat().replace("+00:00", "Z")


def _group_key(creator_id: object, correspondent_email: object) -> tuple[str, str]:
    creator = str(creator_id or "").strip()
    email = normalize_address(correspondent_email)
    if not creator or not email:
        raise ValueError("Follow-up decisions require a Creator and one valid correspondent email.")
    return creator, email


def _ensure_preference(connection, creator_id: str, email: str, timestamp: str) -> None:
    connection.execute(
        "INSERT OR IGNORE INTO mail_follow_up_preferences("
        "creator_id,correspondent_email_normalized,created_at,updated_at) VALUES (?,?,?,?)",
        (creator_id, email, timestamp, timestamp),
    )


def set_snooze(factory, creator_id: object, correspondent_email: object, snooze_until: datetime,
               *, now: datetime | None = None) -> None:
    """Persist a future UTC snooze without changing derived Mail state."""
    creator, email = _group_key(creator_id, correspondent_email)
    current = _utc_now(now)
    until = snooze_until.astimezone(timezone.utc).replace(microsecond=0)
    if until <= current:
        raise ValueError("Snooze time must be in the future.")
    timestamp = _utc_text(current)
    with factory.write_transaction() as connection:
        _ensure_preference(connection, creator, email, timestamp)
        connection.execute(
            "UPDATE mail_follow_up_preferences SET snooze_until=?,updated_at=? "
            "WHERE creator_id=? AND correspondent_email_normalized=?",
            (_utc_text(until), timestamp, creator, email),
        )


def clear_snooze(factory, creator_id: object, correspondent_email: object, *, now: datetime | None = None) -> None:
    creator, email = _group_key(creator_id, correspondent_email)
    timestamp = _utc_text(_utc_now(now))
    with factory.write_transaction() as connection:
        _ensure_preference(connection, creator, email, timestamp)
        connection.execute(
            "UPDATE mail_follow_up_preferences SET snooze_until=NULL,updated_at=? "
            "WHERE creator_id=? AND correspondent_email_normalized=?",
            (timestamp, creator, email),
        )


def stop_follow_up(factory, creator_id: object, correspondent_email: object, *, now: datetime | None = None) -> None:
    creator, email = _group_key(creator_id, correspondent_email)
    timestamp = _utc_text(_utc_now(now))
    with factory.write_transaction() as connection:
        _ensure_preference(connection, creator, email, timestamp)
        connection.execute(
            "UPDATE mail_follow_up_preferences "
            "SET stopped_at=COALESCE(stopped_at,?),updated_at=? "
            "WHERE creator_id=? AND correspondent_email_normalized=?",
            (timestamp, timestamp, creator, email),
        )


def resume_follow_up(factory, creator_id: object, correspondent_email: object, *, now: datetime | None = None) -> None:
    creator, email = _group_key(creator_id, correspondent_email)
    timestamp = _utc_text(_utc_now(now))
    with factory.write_transaction() as connection:
        _ensure_preference(connection, creator, email, timestamp)
        connection.execute(
            "UPDATE mail_follow_up_preferences SET stopped_at=NULL,updated_at=? "
            "WHERE creator_id=? AND correspondent_email_normalized=?",
            (timestamp, creator, email),
        )


def add_to_export_queue(factory, creator_id: object, correspondent_email: object,
                        *, now: datetime | None = None) -> None:
    """Add one group idempotently while retaining its original decision time."""
    creator, email = _group_key(creator_id, correspondent_email)
    timestamp = _utc_text(_utc_now(now))
    with factory.write_transaction() as connection:
        _ensure_preference(connection, creator, email, timestamp)
        connection.execute(
            "UPDATE mail_follow_up_preferences "
            "SET export_queue_added_at=COALESCE(export_queue_added_at,?),updated_at=? "
            "WHERE creator_id=? AND correspondent_email_normalized=?",
            (timestamp, timestamp, creator, email),
        )


def remove_from_export_queue(factory, creator_id: object, correspondent_email: object,
                             *, now: datetime | None = None) -> bool:
    """Remove only explicit membership; keeping the row preserves other decisions."""
    creator, email = _group_key(creator_id, correspondent_email)
    timestamp = _utc_text(_utc_now(now))
    with factory.write_transaction() as connection:
        result = connection.execute(
            "UPDATE mail_follow_up_preferences SET export_queue_added_at=NULL,updated_at=? "
            "WHERE creator_id=? AND correspondent_email_normalized=? AND export_queue_added_at IS NOT NULL",
            (timestamp, creator, email),
        )
    return result.rowcount == 1


def _preferences(factory) -> dict[tuple[str, str], dict]:
    with factory.read_connection() as connection:
        rows = connection.execute("SELECT * FROM mail_follow_up_preferences").fetchall()
    return {(str(row["creator_id"]), str(row["correspondent_email_normalized"])): dict(row) for row in rows}


def follow_up_state(factory, *, now: datetime | None = None) -> list[dict]:
    """Annotate the current Gate 4 read model without changing its derivation."""
    current = _utc_now(now)
    preferences = _preferences(factory)
    rows = []
    for item in derive_follow_up_queue(factory, now=current):
        result = dict(item)
        preference = preferences.get((item["creator_id"], item["correspondent_email"]), {})
        snooze_until = _parse_utc(preference.get("snooze_until"))
        stopped_at = preference.get("stopped_at")
        if stopped_at:
            actionability = "stopped"
        elif snooze_until is not None and current < snooze_until:
            actionability = "snoozed"
        else:
            actionability = "normal"
        result.update({
            "snooze_until": preference.get("snooze_until"),
            "stopped_at": stopped_at,
            "export_queue_added_at": preference.get("export_queue_added_at"),
            "actionability": actionability,
        })
        rows.append(result)
    return rows


def actionable_follow_up_queue(factory, *, now: datetime | None = None) -> list[dict]:
    return [item for item in follow_up_state(factory, now=now) if item["actionability"] == "normal"]


def export_queue(factory, *, now: datetime | None = None) -> list[dict]:
    """Return only queued groups that still have current derived Mail evidence."""
    queued = [item for item in follow_up_state(factory, now=now) if item["export_queue_added_at"]]
    if not queued:
        return []
    creator_ids = sorted({item["creator_id"] for item in queued})
    placeholders = ",".join("?" for _creator_id in creator_ids)
    with factory.read_connection() as connection:
        names = {
            str(row["creator_id"]): str(row["name"] or "")
            for row in connection.execute(
                f"SELECT creator_id,name FROM creators WHERE creator_id IN ({placeholders})", creator_ids
            )
        }
    result = []
    for item in queued:
        row = {key: item.get(key) for key in EXPORT_HEADERS}
        row["creator_name"] = names.get(item["creator_id"], "")
        result.append(row)
    return result


def write_export_csv(factory, destination: Path, *, now: datetime | None = None) -> Path:
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=EXPORT_HEADERS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(export_queue(factory, now=now))
    return path


def write_export_xlsx(factory, destination: Path, *, now: datetime | None = None) -> Path:
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    try:
        sheet = workbook.active
        sheet.title = "Follow Up"
        sheet.append(list(EXPORT_HEADERS))
        for item in export_queue(factory, now=now):
            sheet.append([item.get(header) for header in EXPORT_HEADERS])
        workbook.save(path)
    finally:
        workbook.close()
    return path


def reconcile_export_queue_from_sent(factory, mail_message_ids: Iterable[object],
                                     *, now: datetime | None = None) -> int:
    """Clear only queue memberships proven to have a later, reliable Sent fact.

    Both chronology and local-observation timestamps must meet the persisted
    export decision boundary. This deliberately retains ambiguous, malformed,
    and old-backfill facts.
    """
    ids = sorted({str(value) for value in mail_message_ids if str(value or "").strip()})
    if not ids:
        return 0
    placeholders = ",".join("?" for _identifier in ids)
    statuses = ",".join("?" for _status in RELIABLE_MATCH_STATUSES)
    query = f"""
        SELECT DISTINCT p.creator_id,p.correspondent_email_normalized,p.export_queue_added_at,
               m.message_at,m.observed_at
        FROM mail_follow_up_preferences p
        JOIN mail_messages m ON m.mail_message_id IN ({placeholders})
        JOIN mail_message_addresses a ON a.mail_message_id=m.mail_message_id
        WHERE p.export_queue_added_at IS NOT NULL
          AND m.direction='outbound'
          AND m.match_status IN ({statuses})
          AND m.matched_creator_id=a.matched_creator_id
          AND a.role='to'
          AND a.match_status IN ({statuses})
          AND p.creator_id=a.matched_creator_id
          AND p.correspondent_email_normalized=a.normalized_address
    """
    with factory.read_connection() as connection:
        candidates = [dict(row) for row in connection.execute(query, (*ids, *RELIABLE_MATCH_STATUSES, *RELIABLE_MATCH_STATUSES))]
    cleared = 0
    timestamp = _utc_text(_utc_now(now))
    with factory.write_transaction() as connection:
        for candidate in candidates:
            decision_at = _parse_utc(candidate["export_queue_added_at"])
            message_at = _parse_utc(candidate["message_at"])
            observed_at = _parse_utc(candidate["observed_at"])
            if not decision_at or not message_at or not observed_at:
                continue
            if message_at < decision_at or observed_at < decision_at:
                continue
            result = connection.execute(
                "UPDATE mail_follow_up_preferences SET export_queue_added_at=NULL,updated_at=? "
                "WHERE creator_id=? AND correspondent_email_normalized=? AND export_queue_added_at=?",
                (timestamp, candidate["creator_id"], candidate["correspondent_email_normalized"], candidate["export_queue_added_at"]),
            )
            cleared += result.rowcount
    return cleared
